"""
This is a program to web scrape match data for FIRST Global from the results webpage.
https://results.first.global/

To set this up:
1. You must be able to run this python script in some IDE.
2. Install ChromeDriver (https://developer.chrome.com/docs/chromedriver).
3. Change the string `chrome_driver_path` to be the path to your ChromeDriver.
4. Set up an Excel workbook.
5. Change the string `excel_wb_path` to be the path to your Excel file.
6. Change the string `excel_wb_tab_name` to be the tab in your Excel file where you put the matches.

When you run the script, it should insert the match data into the tab of your excel file. This will
need to be run manually every time you want new data.

If you have any questions, contact Mackensie at mackensiekim24@gmail.com
"""

# import modules

from selenium import webdriver
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.common.by import By
from selenium.common.exceptions import TimeoutException
from openpyxl import load_workbook

# constants to change!
chrome_driver_path = "/Users/user1/Desktop/chromedriver2"
excel_wb_path = "/Users/user1/Desktop/globalexcelfile.xlsx"
excel_wb_tab_name = "Sheet1"

# get website
options = webdriver.ChromeOptions()
options.add_experimental_option('excludeSwitches', ['enable-logging'])
driver = webdriver.Chrome(chrome_driver_path, options=options) # CHANGE THIS
driver.get("https://results.first.global/")

wait = WebDriverWait(driver,10)

# press button to get to matches
path = "/html/body/div/div/div/div[2]/div/div/div/div/button[2]"
button_element = wait.until(EC.presence_of_element_located((By.XPATH,path)))
button_element.click()
print("button has been clicked")

# get actual data
wrapper = "/html/body/div/div/div/div[2]/div[3]/div/div/div/" # gets to whole ranking matches block

match_number = wait.until(EC.presence_of_all_elements_located((By.XPATH, wrapper+"div/p")))
red_1 = wait.until(EC.presence_of_all_elements_located((By.XPATH, wrapper+"div/div/div/div[1]/div[1]/a/span")))
red_2 = wait.until(EC.presence_of_all_elements_located((By.XPATH, wrapper+"div/div/div/div[1]/div[2]/a/span")))
red_3 = wait.until(EC.presence_of_all_elements_located((By.XPATH, wrapper+"div/div/div/div[1]/div[3]/a/span")))
blue_1 = wait.until(EC.presence_of_all_elements_located((By.XPATH, wrapper+"div/div/div/div[2]/div[1]/a/span")))
blue_2 = wait.until(EC.presence_of_all_elements_located((By.XPATH, wrapper+"div/div/div/div[2]/div[2]/a/span")))
blue_3 = wait.until(EC.presence_of_all_elements_located((By.XPATH, wrapper+"div/div/div/div[2]/div[3]/a/span")))
red_score = wait.until(EC.presence_of_all_elements_located((By.XPATH, wrapper+"div/div[2]/div[1]")))
blue_score = wait.until(EC.presence_of_all_elements_located((By.XPATH, wrapper+"div/div[2]/div[2]")))
match_time = wait.until(EC.presence_of_all_elements_located((By.XPATH, wrapper+"div/div[2]")))

# Load excel workbook
wbpath = excel_wb_path
wb = load_workbook(wbpath)
ws = wb[excel_wb_tab_name]

print("workbook loaded")

# write in excel workbook
completed_matches = [None]*len(match_number)

for i in range(len(match_number)):
    ws.cell(2 + i, 2).value = match_number[i].text
    ws.cell(2 + i, 3).value = red_1[i].text
    ws.cell(2 + i, 4).value = red_2[i].text
    ws.cell(2 + i, 5).value = red_3[i].text
    ws.cell(2 + i, 6).value = blue_1[i].text
    ws.cell(2 + i, 7).value = blue_2[i].text
    ws.cell(2 + i, 8).value = blue_3[i].text
    if match_number[i].get_dom_attribute("class") == "MuiTypography-root MuiTypography-body1 css-17mrqxr":
        completed_matches[i] = False
    else:
        completed_matches[i] = True

print("completed matches is done")

# get the predicted match time or scores
completed_match_counter = 0

for i in range(len(match_number)):
    if completed_matches[i]:
        ws.cell(2 + i, 9).value = red_score[completed_match_counter].text
        ws.cell(2 + i, 10).value = blue_score[completed_match_counter].text
        completed_match_counter += 1
    else:
        ws.cell(2 + i, 11).value = match_time[i].text

# end stuff
driver.close()
wb.save(wbpath)
print("writing finished")
